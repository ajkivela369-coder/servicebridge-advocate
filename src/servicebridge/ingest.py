from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from .models import EvidenceChunk, EvidenceClass, SourceDocument
from .redaction import redact_text
from .text import chunk_text, normalize_whitespace


@dataclass(slots=True)
class ExtractedDocument:
    text: str
    page_count: int | None = None


Extractor = Callable[[Path], ExtractedDocument]


def _extract_text(path: Path) -> ExtractedDocument:
    return ExtractedDocument(path.read_text(encoding="utf-8", errors="replace"))


def _extract_json(path: Path) -> ExtractedDocument:
    data = json.loads(path.read_text(encoding="utf-8"))
    return ExtractedDocument(json.dumps(data, indent=2, ensure_ascii=False))


def _extract_pdf(path: Path) -> ExtractedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF support requires: pip install 'servicebridge-advocate[documents]'") from exc
    reader = PdfReader(path)
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append(f"[Page {index}]\n{page.extract_text() or ''}")
    return ExtractedDocument("\n\n".join(pages), page_count=len(reader.pages))


def _extract_docx(path: Path) -> ExtractedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX support requires: pip install 'servicebridge-advocate[documents]'") from exc
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text for cell in row.cells))
    return ExtractedDocument("\n\n".join(blocks))


EXTRACTORS: dict[str, Extractor] = {
    ".txt": _extract_text,
    ".md": _extract_text,
    ".json": _extract_json,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
}


def ingest_file(
    path: str | Path,
    evidence_class: EvidenceClass = EvidenceClass.UNKNOWN,
    *,
    redact: bool = True,
    title: str | None = None,
    document_date: str | None = None,
    author: str | None = None,
) -> tuple[SourceDocument, list[EvidenceChunk]]:
    file_path = Path(path).resolve()
    if not file_path.is_file():
        raise FileNotFoundError(file_path)
    extractor = EXTRACTORS.get(file_path.suffix.lower())
    if extractor is None:
        raise ValueError(f"Unsupported file type: {file_path.suffix or '(none)'}")

    raw = file_path.read_bytes()
    sha256 = hashlib.sha256(raw).hexdigest()
    source_id = f"src_{sha256[:16]}"
    extracted = extractor(file_path)
    text = normalize_whitespace(extracted.text)
    counts: dict[str, int] = {}
    if redact:
        text, counts = redact_text(text)

    source = SourceDocument(
        source_id=source_id,
        title=title or file_path.stem,
        path=str(file_path),
        evidence_class=evidence_class,
        document_date=document_date,
        author=author,
        page_count=extracted.page_count,
        sha256=sha256,
        redacted=redact,
        metadata={"redaction_counts": counts, "original_suffix": file_path.suffix.lower()},
    )
    chunks = [
        EvidenceChunk(
            chunk_id=f"{source_id}_c{ordinal:04d}",
            source_id=source_id,
            text=chunk,
            ordinal=ordinal,
            locator=_infer_locator(chunk, ordinal),
            evidence_class=evidence_class,
        )
        for ordinal, chunk in enumerate(chunk_text(text), start=1)
    ]
    return source, chunks


def _infer_locator(text: str, ordinal: int) -> str:
    first_line = text.splitlines()[0].strip() if text else ""
    if first_line.startswith("[Page ") and first_line.endswith("]"):
        return first_line.strip("[]")
    return f"chunk {ordinal}"
